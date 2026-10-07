import re
import os


def clean_text(text: str) -> str:
    """
    Clean extracted resume text while preserving useful
    line structure for name/section detection.
    """

    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    # Remove excessive spaces/tabs
    text = re.sub(r'[ \t]+', ' ', text)

    # Remove excessive blank lines
    text = re.sub(r'\n{3,}', '\n\n', text)

    # Remove spaces around newlines
    text = re.sub(r' *\n *', '\n', text)

    return text.strip()


# ============================================================
# PDF
# ============================================================

def extract_text_from_pdf(file) -> str:
    import pdfplumber

    pages_text = []

    with pdfplumber.open(file) as pdf:

        for page in pdf.pages:

            text = page.extract_text(
                x_tolerance=2,
                y_tolerance=3
            )

            if text:
                pages_text.append(text)

    if not pages_text:
        return ""

    return clean_text("\n".join(pages_text))


# ============================================================
# DOCX
# ============================================================

def extract_text_from_docx(file) -> str:
    from docx import Document

    doc = Document(file)

    parts = []

    # --------------------------------------------------------
    # Normal paragraphs
    # --------------------------------------------------------

    for paragraph in doc.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------

    for table in doc.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                parts.append(" | ".join(row_text))

    if not parts:
        return ""

    return clean_text("\n".join(parts))


# ============================================================
# MAIN PARSER
# ============================================================

def parse_resume(file) -> str:

    filename = file.name.lower()

    if filename.endswith(".pdf"):

        return extract_text_from_pdf(file)

    elif filename.endswith(".docx"):

        return extract_text_from_docx(file)

    else:

        raise ValueError(
            f"Unsupported file type: '{file.name}'. "
            "Please upload a PDF or DOCX file."
        )


# ============================================================
# TEST
# ============================================================

if __name__ == "__main__":

    import sys

    BASE_DIR = os.path.dirname(
        os.path.abspath(__file__)
    )

    ROOT = os.path.abspath(
        os.path.join(BASE_DIR, "..")
    )

    sys.path.insert(0, ROOT)

    pdf_path = os.path.join(
        ROOT,
        "sample_resume.pdf"
    )

    text = extract_text_from_pdf(pdf_path)

    print(text[:1000])

    print(
        f"\nTotal characters: {len(text)}"
    )